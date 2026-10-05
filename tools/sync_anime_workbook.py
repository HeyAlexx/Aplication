from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_BASELINE = PROJECT_DIR.parents[1] / "Referencias" / "Anime Data V1.xlsx"
ANIME_PATH = PROJECT_DIR / "data" / "anime.json"

QUARTERS = {1: "Winter", 2: "Spring", 3: "Summer", 4: "Fall"}
FORMAT_MAP = {"Serie": "Serie", "Movie": "Película", "OVA": "OVA", "ONA": "ONA", "Especial": "Especial"}
PRODUCTION_MAP = {"Emision": "Emisión", "Emisión": "Emisión", "Finalizado": "Finalizado", "Pausado": "Pausado"}
VIEWING_MAP = {"Visto": "Visto", "No visto": "No Visto", "No Visto": "No Visto", "Por Ver": "Por Ver", "Ignorado": "Ignorado"}
MANAGED_FIELDS = {
    1: "sourceCode",
    2: "title",
    3: "productionStatus",
    4: "format",
    5: "chapters",
    6: "seasonsCount",
    7: "activeSeason",
    8: "viewingStatus",
    9: "emissionYear",
    10: "emissionSeason",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fusiona cambios de un Excel de anime sin perder metadatos enriquecidos.")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args()


def text(value: object) -> str:
    return str(value).strip() if value not in (None, "") else ""


def code_text(value: object) -> str:
    if isinstance(value, (int, float)) and value == int(value):
        return str(int(value))
    return text(value)


def integer(value: object, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def numeric_text(value: object) -> str:
    if isinstance(value, (int, float)) and value == int(value):
        return str(int(value))
    return text(value)


def raw_row(row_number: int, row: tuple[object, ...]) -> dict[str, object]:
    values = list(row[:10]) + [None] * max(0, 10 - len(row))
    code, title, status, content_type, chapters, seasons, active, watching, year, quarter = values[:10]
    return {
        "row": row_number,
        "raw": [code_text(code), text(title), text(status), text(content_type), numeric_text(chapters), numeric_text(seasons), text(active), text(watching), numeric_text(year), numeric_text(quarter)],
        "sourceCode": code_text(code),
        "title": text(title),
        "productionStatus": PRODUCTION_MAP.get(text(status), text(status) or "Pendiente"),
        "format": FORMAT_MAP.get(text(content_type), text(content_type) or "Serie"),
        "chapters": integer(chapters),
        "seasonsCount": integer(seasons),
        "activeSeason": text(active),
        "viewingStatus": VIEWING_MAP.get(text(watching), text(watching) or "Por Ver"),
        "emissionYear": integer(year) or None,
        "emissionSeason": QUARTERS.get(integer(quarter), ""),
    }


def read_catalog(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Catalogo"]
        result = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True), start=2):
            item = raw_row(row_number, row)
            if item["sourceCode"] and item["title"]:
                result.append(item)
        return result
    finally:
        workbook.close()


def build_index(items: list[dict[str, object]], code_key: str = "sourceCode") -> tuple[dict[str, list[dict[str, object]]], dict[str, list[dict[str, object]]]]:
    by_code: dict[str, list[dict[str, object]]] = defaultdict(list)
    by_title: dict[str, list[dict[str, object]]] = defaultdict(list)
    for item in items:
        by_code[text(item.get(code_key))].append(item)
        by_title[text(item.get("title")).casefold()].append(item)
    return by_code, by_title


def make_new_record(item: dict[str, object], timestamp: str) -> dict[str, object]:
    record = {
        "id": f"anime-sheet-{item['sourceCode']}",
        "sourceCode": item["sourceCode"],
        "sourceRow": item["row"],
        "title": item["title"],
        "type": "anime",
        "categoryGeneral": "Anime",
        "format": item["format"],
        "productionStatus": item["productionStatus"],
        "viewingStatus": item["viewingStatus"],
        "genre": "Anime",
        "year": str(item["emissionYear"] or ""),
        "emissionYear": item["emissionYear"],
        "emissionSeason": item["emissionSeason"],
        "chapters": item["chapters"],
        "seasonsCount": item["seasonsCount"],
        "activeSeason": item["activeSeason"],
        "tags": [item["format"], item["viewingStatus"], item["productionStatus"]],
        "image": "",
        "description": "Registro agregado desde la hoja de anime de Altoidss.",
        "featured": False,
        "status": "Activo",
        "source": "google-sheet-anime",
        "sheetUpdatedAt": timestamp,
    }
    return record


def sheet_values(record: dict[str, object]) -> list[object]:
    quarter = next((number for number, season in QUARTERS.items() if season == record.get("emissionSeason")), "")
    format_value = "Movie" if record.get("format") == "Película" else record.get("format", "")
    status_value = "Emision" if record.get("productionStatus") == "Emisión" else record.get("productionStatus", "")
    watching_value = "No visto" if record.get("viewingStatus") == "No Visto" else record.get("viewingStatus", "")
    return [
        record.get("sourceCode", ""), record.get("title", ""), status_value, format_value,
        record.get("chapters", 0), record.get("seasonsCount", 0), record.get("activeSeason", ""),
        watching_value, record.get("emissionYear") or "", quarter,
    ]


def main() -> None:
    args = parse_args()
    baseline = read_catalog(args.baseline)
    incoming = read_catalog(args.workbook)
    records = json.loads(ANIME_PATH.read_text(encoding="utf-8"))
    baseline_by_code, baseline_by_title = build_index(baseline)
    records_by_code, records_by_title = build_index(records)
    timestamp = datetime.now().astimezone().isoformat()
    changed = []
    added = []

    for item in incoming:
        baseline_matches = baseline_by_title[item["title"].casefold()]
        baseline_item = baseline_matches[0] if len(baseline_matches) == 1 else None
        if baseline_item is None and len(baseline_by_code[item["sourceCode"]]) == 1:
            baseline_item = baseline_by_code[item["sourceCode"]][0]

        record_matches = records_by_title[item["title"].casefold()]
        record = record_matches[0] if len(record_matches) == 1 else None
        if record is None and len(records_by_code[item["sourceCode"]]) == 1:
            record = records_by_code[item["sourceCode"]][0]

        if baseline_item is None and record is None:
            record = make_new_record(item, timestamp)
            records.append(record)
            records_by_code[item["sourceCode"]].append(record)
            records_by_title[item["title"].casefold()].append(record)
            added.append(record)
            continue

        if baseline_item is None or record is None:
            continue

        updated_fields = []
        for column, key in MANAGED_FIELDS.items():
            if item["raw"][column - 1] == baseline_item["raw"][column - 1]:
                continue
            new_value = item[key]
            if record.get(key) != new_value:
                record[key] = new_value
                updated_fields.append(key)

        if updated_fields:
            record["year"] = str(record.get("emissionYear") or "")
            record["tags"] = [record.get("format", ""), record.get("viewingStatus", ""), record.get("productionStatus", "")]
            record["sheetUpdatedAt"] = timestamp
            changed.append({"id": record.get("id"), "sourceRow": record.get("sourceRow"), "fields": updated_fields})

    if args.apply:
        ANIME_PATH.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    changed_ids = {item["id"] for item in changed}
    sheet_updates = [
        {"id": record.get("id"), "targetRow": record.get("sourceRow"), "values": sheet_values(record)}
        for record in records if record.get("id") in changed_ids
    ]
    print(json.dumps({
        "applied": args.apply,
        "changedCount": len(changed),
        "addedCount": len(added),
        "total": len(records),
        "changed": changed,
        "added": [{"id": item["id"], "sourceCode": item["sourceCode"], "title": item["title"], "values": sheet_values(item)} for item in added],
        "sheetUpdates": sheet_updates,
        "timestamp": timestamp,
    }, ensure_ascii=True))


if __name__ == "__main__":
    main()
